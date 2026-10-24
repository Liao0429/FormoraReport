"""结果预览页：读取生成的 DOCX 文本并展示；支持导出 PDF。"""
from __future__ import annotations

from pathlib import Path

from PySide6.QtCore import QThread, Signal
from PySide6.QtWidgets import QHBoxLayout, QLabel, QMessageBox, QPlainTextEdit, QVBoxLayout, QWidget

from formora.core.errors import FormoraError
from formora.ui.widgets import Card, primary_button, secondary_button


class _ReviewThread(QThread):
    """后台执行 AI 审查，避免阻塞 UI 主线程。"""

    result_ready = Signal(object)
    error_raised = Signal(str)

    def __init__(self, fn, *args, **kwargs):
        super().__init__()
        self._fn = fn
        self._args = args
        self._kwargs = kwargs

    def run(self):
        try:
            self.result_ready.emit(self._fn(*self._args, **self._kwargs))
        except Exception as e:
            self.error_raised.emit(str(e))


class PreviewPage(QWidget):
    def __init__(self, state, navigate):
        super().__init__()
        self.state = state
        self.navigate = navigate
        self._review_thread: _ReviewThread | None = None
        self._build()

    def _build(self):
        lay = QVBoxLayout(self)
        lay.setContentsMargins(24, 24, 24, 24)
        lay.setSpacing(14)

        title = QLabel("结果预览")
        title.setObjectName("AppTitle")
        lay.addWidget(title)

        card = Card("报告内容预览")
        self.text_edit = QPlainTextEdit()
        self.text_edit.setReadOnly(True)
        self.text_edit.setPlaceholderText("生成后可在此预览文本内容")
        card.add_widget(self.text_edit)
        lay.addWidget(card, 1)

        ops = QHBoxLayout()
        back_btn = secondary_button("← 返回生成")
        back_btn.clicked.connect(lambda: self.navigate("generate"))
        self.ai_review_btn = secondary_button("✨ AI 审查")
        self.ai_review_btn.clicked.connect(self._ai_review)
        export_btn = primary_button("导出 PDF")
        export_btn.clicked.connect(self._export_pdf)
        ops.addWidget(back_btn)
        ops.addWidget(self.ai_review_btn)
        ops.addStretch(1)
        ops.addWidget(export_btn)
        lay.addLayout(ops)

    def _ai_sources_text(self) -> str:
        """把当前映射中的 AI 来源追踪信息格式化为文本。"""
        sources = [m for m in self.state.mappings if m.source or m.needs_confirmation]
        if not sources:
            return ""
        lines = ["\n【AI 来源追踪】"]
        for m in sources:
            flag = "（待确认）" if m.needs_confirmation else ""
            lines.append(f"- {m.template_field}{flag}: {m.source or '无依据'}")
        return "\n".join(lines)

    def _ai_review(self):
        from formora.ai import client

        if not self.state.output_path or not Path(self.state.output_path).exists():
            QMessageBox.warning(self, "提示", "还没有可审查的输出文件。")
            return
        cfg = self.state.config
        if not client.is_configured(cfg):
            QMessageBox.information(
                self, "AI 模式未启用",
                "AI 审查需要在「设置 → AI 智能模式」中配置 API 地址、模型与密钥。")
            return
        # 隐私确认：AI 审查会把整份报告文本发送到所配置的 API
        if not cfg.get("privacy", {}).get("review_warning_shown"):
            ret = QMessageBox.question(
                self, "隐私提示",
                "AI 审查会把当前报告的完整文本（含已填入的数据）发送到"
                "你配置的模型 API。\n请确认你信任该 API 服务商。是否继续？")
            if ret != QMessageBox.StandardButton.Yes:
                return
            cfg.set("privacy", {**cfg.get("privacy", {}), "review_warning_shown": True})
            cfg.save()
        try:
            from docx import Document
            doc = Document(self.state.output_path)
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for i, table in enumerate(doc.tables):
                parts.append(f"[表格 {i+1}]")
                for row in table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            report_text = "\n".join(parts)
        except Exception as e:
            QMessageBox.warning(self, "读取报告失败", str(e))
            return
        # 后台线程化：AI 请求可能耗时，不阻塞 UI
        self._set_review_busy(True)
        self.text_edit.append("\n\n⏳ AI 审查中，请稍候…")
        self._review_thread = _ReviewThread(client.review_report, cfg, report_text)
        self._review_thread.result_ready.connect(self._on_review_done)
        self._review_thread.error_raised.connect(self._on_review_error)
        self._review_thread.start()

    def _set_review_busy(self, busy: bool):
        if hasattr(self, "ai_review_btn"):
            self.ai_review_btn.setEnabled(not busy)

    def _on_review_done(self, issues):
        self._set_review_busy(False)
        self.refresh()  # 恢复原始报告预览
        if not issues:
            self.text_edit.append("\n【AI 审查】未发现问题。（AI 判断仅供参考，请人工复核）")
            return
        lines = ["\n【AI 审查】发现以下问题："]
        for iss in issues:
            sev = iss.get("severity", "info")
            icon = "🔴" if sev == "error" else ("⚠️" if sev == "warning" else "ℹ️")
            lines.append(f"{icon} {iss.get('field', '')}: {iss.get('issue', '')}")
        self.text_edit.append("\n".join(lines))
        self.text_edit.append("\n（AI 判断仅供参考，请人工复核）")

    def _on_review_error(self, msg):
        self._set_review_busy(False)
        self.refresh()
        self.text_edit.append(f"\n【AI 审查失败】{msg}")

    def refresh(self):
        if not self.state.output_path or not Path(self.state.output_path).exists():
            self.text_edit.setPlainText("（还没有可预览的输出文件）")
            return
        try:
            from docx import Document
            doc = Document(self.state.output_path)
            parts = []
            for p in doc.paragraphs:
                if p.text.strip():
                    parts.append(p.text)
            for i, table in enumerate(doc.tables):
                parts.append(f"[表格 {i+1}]")
                for row in table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            # 提示图片数量（文本预览无法展示图片，引导用户用 Word 查看排版）
            if doc.inline_shapes:
                parts.append(f"\n📷 报告包含 {len(doc.inline_shapes)} 张图片（文本预览不显示图片，请用 Word 打开查看排版效果）。")
            parts.append(self._ai_sources_text())
            self.text_edit.setPlainText("\n".join(parts))
        except Exception as e:
            self.text_edit.setPlainText(f"读取预览失败：{e}")

    def _export_pdf(self):
        if not self.state.output_path:
            QMessageBox.warning(self, "提示", "还没有输出文件。")
            return
        from formora.core import pdf_exporter
        try:
            out = self.state.output_path.replace(".docx", ".pdf")
            pdf_exporter.export_pdf(self.state.output_path, out)
            QMessageBox.information(self, "导出成功", f"PDF 已导出：\n{out}")
        except FormoraError as e:
            QMessageBox.warning(self, "导出失败", str(e))