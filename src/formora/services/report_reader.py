"""报告文本读取服务：从生成的 DOCX 中抽取可读文本。

把原本写在 preview_page 里的文本抽取逻辑抽成纯函数，
供桌面端 UI 预览、AI 审查、以及未来的后端 API 复用。
"""
from __future__ import annotations

from pathlib import Path


def extract_report_text(docx_path: str | Path) -> str:
    """从生成的 DOCX 中抽取段落+表格文本，供预览与 AI 审查。

    - 段落文本
    - 表格文本（以 [表格 N] 标注）
    - 图片数量提示（文本预览无法展示图片）
    """
    from docx import Document

    docx_path = Path(docx_path)
    doc = Document(str(docx_path))
    parts: list[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    for i, table in enumerate(doc.tables):
        parts.append(f"[表格 {i + 1}]")
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    if doc.inline_shapes:
        parts.append(
            f"\n📷 报告包含 {len(doc.inline_shapes)} 张图片"
            "（文本预览不显示图片，请用 Word 打开查看排版效果）。")
    return "\n".join(parts)


def extract_report_text_with_sources(docx_path: str | Path, ai_sources_text: str = "") -> str:
    """抽取报告文本并追加 AI 来源追踪。"""
    text = extract_report_text(docx_path)
    if ai_sources_text:
        text += "\n" + ai_sources_text
    return text