"""DOCX 模板解析：检测占位符、循环块、章节、表格、图片、页眉页脚。

仅支持 .docx（OOXML）。旧版 .doc 不支持，给出明确错误。
"""

from __future__ import annotations

import re
from pathlib import Path

from .errors import DocumentFormatError, ExternalDependencyError, UserError
from .models import TableInfo, TemplateInfo

# 占位符：{{ 字段名 }}。允许中英文、数字、空格、点、连字符、括号（如 {{学生 姓名}}、{{chart-1}}、{{实验名称(一)}}）。
# 排除换行与花括号本身，保持 docxtpl/jinja 兼容。
_PLACEHOLDER_RE = re.compile(r"\{\{\s*([^{}\n]+?)\s*\}\}")
# 其它常见占位符写法：${字段名} / 【字段名】（Word 域/中文提示习惯）
_DOLLAR_RE = re.compile(r"\$\{\s*([^{}\n]+?)\s*\}")
_CN_BRACKET_RE = re.compile(r"【\s*([^】\n]+?)\s*】")


def _find_placeholders(text: str) -> set[str]:
    """检测文本中的占位符字段名，兼容 {{}} / ${} / 【】 三种写法。"""
    found: set[str] = set()
    found.update(_PLACEHOLDER_RE.findall(text))
    found.update(_DOLLAR_RE.findall(text))
    found.update(_CN_BRACKET_RE.findall(text))
    return found


_LOOP_RE = re.compile(r"\{\%\s*for\s+(\w+)\s+in\s+(\w+)\s*\%\}")
# 循环体内字段引用：{{ row['项目'] }} / {{ row.项目 }} / {{ item["name"] }}
_LOOP_FIELD_RE = re.compile(
    r"\{\{\s*\w+\s*\.?\s*\[?\s*['\"]?([\w\u4e00-\u9fa5]+)['\"]?\s*\]?\s*\}\}"
)
# 循环块整体：{% for x in y %} ... {% endfor %}（用于从文本中剔除循环体内引用）
_LOOP_BODY_RE = re.compile(
    r"\{\%\s*for\s+\w+\s+in\s+\w+\s*\%\}.*?\{\%\s*endfor\s*\%\}", re.DOTALL
)

# 章节关键字 -> 规范化章节名
_SECTION_KEYWORDS = [
    ("实验名称", ["实验名称", "实验题目", "题目"]),
    ("实验目的", ["实验目的", "目的"]),
    ("实验环境", ["实验环境", "实验条件", "环境", "软硬件环境"]),
    ("实验步骤", ["实验步骤", "步骤", "操作过程", "实验过程"]),
    ("实验数据", ["实验数据", "数据记录", "原始数据"]),
    ("实验结果", ["实验结果", "结果分析", "结果"]),
    ("问题分析", ["问题分析", "问题讨论", "思考题", "实验问题"]),
    ("实验结论", ["实验结论", "结论", "总结", "心得体会"]),
    ("附录", ["附录"]),
]


def _extract_all_text(element) -> list[str]:
    """递归提取 OOXML 元素下所有 w:t 文本节点的非空文本。

    用于扫描页眉/页脚/表格/嵌套表格中的文本（python-docx 的 paragraphs
    不包含表格单元格，必须遍历底层 XML 才能覆盖完整）。
    """
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    texts: list[str] = []
    for t in element.iter():
        if t.tag == ns + "t" and t.text and t.text.strip():
            texts.append(t.text.strip())
    return texts


def _detect_sections(paragraph_texts: list[str]) -> dict[str, str]:
    sections: dict[str, str] = {}
    for para in paragraph_texts:
        text = para.strip()
        if not text:
            continue
        for canonical, keywords in _SECTION_KEYWORDS:
            if canonical in sections:
                continue
            for kw in keywords:
                idx = text.find(kw)
                # 标题形式：「实验名称」「一、实验名称」「(1) 实验目的」等：
                # 关键字出现在前 8 个字符内，且整段较短（更像标题而非正文）
                if idx != -1 and idx <= 8 and len(text) <= 50:
                    sections[canonical] = text
                    break
    return sections


def parse_template(path: str | Path) -> TemplateInfo:
    path = Path(path)
    if not path.exists():
        raise UserError(f"模板文件不存在: {path}")
    if path.suffix.lower() != ".docx":
        raise DocumentFormatError(
            f"暂不支持 {path.suffix or '无后缀'} 模板，仅支持 .docx 格式。"
            f"请先在 Word 中另存为 .docx 后再导入。"
        )
    try:
        from docx import Document
    except ImportError:
        raise ExternalDependencyError("python-docx 未安装，请先安装依赖。")

    try:
        doc = Document(str(path))
    except Exception as e:
        raise DocumentFormatError(
            f"无法打开 DOCX 模板（文件可能损坏或不是有效 Word 文档）: {e}"
        )

    placeholders: set[str] = set()
    loop_blocks: set[str] = set()
    para_texts: list[str] = []
    tables: list[TableInfo] = []
    has_images = False
    has_header_footer = False

    def scan_paragraph(p):
        nonlocal has_images
        text = p.text or ""
        if text.strip():
            para_texts.append(text)
        # 循环体内的 {{ row['字段'] }} 引用是 docxtpl 循环变量，不是顶层映射占位符，
        # 不应加入 placeholders（否则映射表出现 row['项目'] 这类怪字段）
        loop_body = _LOOP_BODY_RE.sub("", text)
        placeholders.update(_find_placeholders(loop_body))
        for _var, _seq in _LOOP_RE.findall(text):
            loop_blocks.add(_var)
        if p._p.xml.count("<pic:pic") > 0 or "graphic" in p._p.xml:
            has_images = True

    # 正文段落
    for para in doc.paragraphs:
        scan_paragraph(para)

    # 表格
    for ti, table in enumerate(doc.tables):
        header: list[str] = []
        sample: list[list[str]] = []
        for ri, row in enumerate(table.rows[:6]):
            cells = [c.text.strip() for c in row.cells]
            if ri == 0:
                header = cells
            sample.append(cells)
            for cell in row.cells:
                # 循环体剔除：与正文一致，避免表格内 {% for %} 循环体的 row['x'] 误入顶层
                cell_text = cell.text or ""
                cell_body = _LOOP_BODY_RE.sub("", cell_text)
                placeholders.update(_find_placeholders(cell_body))
                cell_paras = [cp.text for cp in cell.paragraphs if cp.text.strip()]
                para_texts.extend(cell_paras)
        tables.append(
            TableInfo(
                index=ti,
                header=header,
                row_count=len(table.rows),
                col_count=len(table.columns),
                sample=sample,
            )
        )

    # 页眉页脚：递归扫描所有 w:t 文本（含页眉页脚里的表格单元格、嵌套表格）
    for section in doc.sections:
        for hf in (section.header, section.footer):
            texts = _extract_all_text(hf._element)
            if texts:
                has_header_footer = True
                para_texts.extend(texts)
                for t in texts:
                    placeholders.update(_find_placeholders(t))

    sections = _detect_sections(para_texts)
    return TemplateInfo(
        path=path,
        placeholders=sorted(placeholders),
        loop_blocks=sorted(loop_blocks),
        sections=sections,
        tables=tables,
        has_images=has_images,
        has_header_footer=has_header_footer,
    )
