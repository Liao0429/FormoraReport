#!/usr/bin/env python3
"""生成真实样例：一个含占位符/循环/图表位的 DOCX 实验报告模板 + CSV 实验数据。

用于 MVP 验证和自动化测试。生成到 tests/samples/。

样例内容完全确定：不依赖随机数、时间戳或当前工作目录，重复运行生成一致结果，
便于在 CI / 自动化测试中复现。
"""

from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

SAMPLES = Path(__file__).resolve().parent.parent / "tests" / "samples"

# 样例 CSV 数据。数值以字符串书写，避免浮点序列化差异，保证字节级可复现。
_SAMPLE_ROWS: list[list[str]] = [
    ["项目", "数值"],
    ["样本A", "12.5"],
    ["样本B", "18.2"],
    ["样本C", "15.7"],
    ["样本D", "21.3"],
]


def _h(doc: Document, text: str) -> None:
    """添加一个加粗段落，作为章节标题。"""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True


def make_template(path: Path) -> None:
    """生成含占位符/循环/图表位的 DOCX 实验报告模板。"""
    doc = Document()
    # 标题
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("实 验 报 告")
    run.bold = True
    run.font.size = Pt(18)

    # 基本信息表
    table = doc.add_table(rows=2, cols=4)
    table.style = "Table Grid"
    cells = [
        ("姓名", "{{姓名}}", "学号", "{{学号}}"),
        ("班级", "{{班级}}", "日期", "{{日期}}"),
    ]
    for ri, row in enumerate(cells):
        for ci, val in enumerate(row):
            table.cell(ri, ci).text = val

    doc.add_paragraph()
    _h(doc, "一、实验名称")
    doc.add_paragraph("{{实验名称}}")
    _h(doc, "二、实验目的")
    doc.add_paragraph("{{实验目的}}")
    _h(doc, "三、实验环境")
    doc.add_paragraph("{{实验环境}}")
    _h(doc, "四、实验步骤")
    doc.add_paragraph("{{实验步骤}}")
    _h(doc, "五、实验数据")
    doc.add_paragraph(
        "{% for row in data_rows %}{{ row['项目'] }}：{{ row['数值'] }}；{% endfor %}"
    )
    _h(doc, "六、实验结果图表")
    doc.add_paragraph("{{ chart }}")
    _h(doc, "七、实验结果")
    doc.add_paragraph("{{实验结果}}")
    _h(doc, "八、问题分析")
    doc.add_paragraph("{{问题分析}}")
    _h(doc, "九、实验结论")
    doc.add_paragraph("{{实验结论}}")

    doc.save(str(path))
    print(f"模板已生成: {path}")


def make_csv(path: Path) -> None:
    """生成 UTF-8（带 BOM）CSV 实验数据。"""
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerows(_SAMPLE_ROWS)
    print(f"数据已生成: {path}")


def main() -> None:
    SAMPLES.mkdir(parents=True, exist_ok=True)
    make_template(SAMPLES / "实验报告模板.docx")
    make_csv(SAMPLES / "实验数据.csv")


if __name__ == "__main__":
    main()
