"""FormoraReport 核心功能自动化测试。

用真实样例文件（tests/samples/）验证：
- 模板解析：占位符、循环块、章节、表格
- 数据读取：CSV
- 字段映射
- 报告生成：输出可打开的 DOCX，内容正确
- 校验
- 错误处理：不支持格式、不存在的文件、损坏模板
"""
import sys
import tempfile
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from formora.core import (
    chart,
    data_reader,
    generator,
    mapper,
    template_parser,
    validator,
)
from formora.core.errors import DocumentFormatError, UserError
from formora.core.models import DataSource, FieldMapping

SAMPLES = Path(__file__).resolve().parent.parent / "tests" / "samples"
# 测试输出写入系统临时目录，避免污染仓库 tests/output 造成杂乱。
OUT = Path(tempfile.mkdtemp(prefix="formora_test_"))
TEMPLATE = SAMPLES / "实验报告模板.docx"
DATA = SAMPLES / "实验数据.csv"


@pytest.fixture(scope="session", autouse=True)
def ensure_samples():
    """确保样例存在（仅缺时生成，避免改动已跟踪的样例文件）。"""
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
    from make_samples import make_csv, make_template
    SAMPLES.mkdir(parents=True, exist_ok=True)
    if not TEMPLATE.exists():
        make_template(TEMPLATE)
    if not DATA.exists():
        make_csv(DATA)
    yield


class TestTemplateParser:
    def test_parse_placeholders(self):
        info = template_parser.parse_template(TEMPLATE)
        assert "实验名称" in info.placeholders
        assert "姓名" in info.placeholders
        assert "学号" in info.placeholders
        assert "chart" in info.placeholders

    def test_parse_loop(self):
        info = template_parser.parse_template(TEMPLATE)
        assert "row" in info.loop_blocks

    def test_parse_sections(self):
        info = template_parser.parse_template(TEMPLATE)
        # 标题为「一、实验名称」等，应被识别
        assert "实验名称" in info.sections
        assert "实验目的" in info.sections
        assert "实验环境" in info.sections
        assert "实验步骤" in info.sections
        assert "实验数据" in info.sections
        assert "实验结果" in info.sections
        assert "问题分析" in info.sections
        assert "实验结论" in info.sections

    def test_parse_tables(self):
        info = template_parser.parse_template(TEMPLATE)
        assert len(info.tables) >= 1
        assert info.tables[0].col_count == 4

    def test_missing_file(self):
        with pytest.raises(UserError):
            template_parser.parse_template(SAMPLES / "不存在.docx")

    def test_unsupported_format(self):
        bad = SAMPLES / "x.doc"
        bad.write_bytes(b"fake")
        try:
            with pytest.raises(DocumentFormatError):
                template_parser.parse_template(bad)
        finally:
            bad.unlink()


class TestDataReader:
    def test_read_csv(self):
        ds = data_reader.read_data(DATA)
        assert ds.kind == "csv"
        assert "项目" in ds.columns
        assert len(ds.rows) == 4

    def test_read_text(self):
        p = SAMPLES / "readme.txt"
        p.write_text("实验说明\n第一行数据", encoding="utf-8")
        try:
            ds = data_reader.read_data(p)
            assert ds.kind == "txt"
            assert "第一行数据" in ds.text
            assert not ds.is_tabular
        finally:
            p.unlink()

    def test_read_markdown(self):
        p = SAMPLES / "note.md"
        p.write_text("# 标题\n正文内容", encoding="utf-8")
        try:
            ds = data_reader.read_data(p)
            assert ds.kind == "md"
            assert "正文内容" in ds.text
        finally:
            p.unlink()

    def test_missing_file(self):
        with pytest.raises(UserError):
            data_reader.read_data(SAMPLES / "不存在.csv")


class TestMapper:
    def test_auto_map_falls_back_to_constant(self):
        info = template_parser.parse_template(TEMPLATE)
        ds = data_reader.read_data(DATA)
        mappings = mapper.auto_map(info, ds)
        # 姓名等不在 CSV 列中 -> constant
        by_field = {m.template_field: m for m in mappings}
        assert by_field["姓名"].source_type == "constant"
        # 项目列应匹配到某占位符？没有直接对应，因此均为 constant
        assert all(m.source_type in ("constant", "column") for m in mappings)

    def test_auto_map_column_match(self):
        # 造一个列名与占位符相同的场景
        info = template_parser.parse_template(TEMPLATE)
        ds = DataSource(path=DATA, kind="csv", columns=["姓名", "学号"], rows=[])
        mappings = mapper.auto_map(info, ds)
        by_field = {m.template_field: m for m in mappings}
        assert by_field["姓名"].source_type == "column"
        assert by_field["姓名"].source_column == "姓名"

    def test_auto_map_alias_match(self):
        """占位符「姓名」应匹配数据列「学生姓名」等别名。"""
        info = template_parser.parse_template(TEMPLATE)
        ds = DataSource(path=DATA, kind="csv", columns=["学生姓名", "学号"], rows=[])
        mappings = mapper.auto_map(info, ds)
        by_field = {m.template_field: m for m in mappings}
        assert by_field["姓名"].source_type == "column"
        assert by_field["姓名"].source_column == "学生姓名"
        assert by_field["学号"].source_column == "学号"

    def test_resolve_values(self):
        mappings = [
            FieldMapping(template_field="姓名", source_type="column", source_column="学生姓名"),
            FieldMapping(template_field="备注", source_type="constant", constant_value="无"),
        ]
        ctx = mapper.resolve_values(mappings, row={"学生姓名": "张三"})
        assert ctx["姓名"] == "张三"
        assert ctx["备注"] == "无"


class TestGenerator:
    def test_generate_docx(self):
        info = template_parser.parse_template(TEMPLATE)
        ds = data_reader.read_data(DATA)
        mappings = mapper.auto_map(info, ds)
        constants = {
            "实验名称": "测试实验",
            "实验目的": "目的",
            "实验环境": "环境",
            "实验步骤": "步骤",
            "实验结果": "结果",
            "问题分析": "问题",
            "实验结论": "结论",
        }
        for m in mappings:
            if m.source_type == "constant" and m.template_field in constants:
                m.constant_value = constants[m.template_field]
        context = mapper.resolve_values(mappings, row=ds.rows[0] if ds.rows else None)
        context["data_rows"] = ds.rows

        # 生成图表
        chart_png = OUT / "test_chart.png"
        chart.generate_line_chart(
            [str(r["项目"]) for r in ds.rows],
            [float(r["数值"]) for r in ds.rows],
            chart_png, title="测试图")
        assert chart_png.exists()

        out = generator.generate_report(
            TEMPLATE, context, OUT / "test_output.docx",
            image_vars={"chart": str(chart_png)})
        assert out.exists()

        # 输出可打开且内容正确
        from docx import Document
        doc = Document(str(out))
        texts = " ".join(p.text for p in doc.paragraphs)
        assert "测试实验" in texts
        assert "样本A" in texts  # 循环表格数据已渲染
        assert "12.5" in texts

    def test_generate_inserts_image(self):
        """验证图片确实作为内嵌对象插入到输出文档。"""
        info = template_parser.parse_template(TEMPLATE)
        ds = data_reader.read_data(DATA)
        mappings = mapper.auto_map(info, ds)
        context = mapper.resolve_values(mappings, row=ds.rows[0] if ds.rows else None)
        context["data_rows"] = ds.rows
        chart_png = OUT / "test_chart2.png"
        chart.generate_line_chart([1, 2, 3], [4, 5, 6], chart_png, title="图")
        out = generator.generate_report(
            TEMPLATE, context, OUT / "test_img.docx",
            image_vars={"chart": str(chart_png)})
        from docx import Document
        doc = Document(str(out))
        assert len(doc.inline_shapes) > 0, "输出文档应包含内嵌图片"

    def test_generate_missing_image_raises(self):
        """绑定的图片文件不存在时应抛出 ExternalDependencyError。"""
        from formora.core.errors import ExternalDependencyError
        info = template_parser.parse_template(TEMPLATE)
        context = {"实验名称": "x"}
        with pytest.raises(ExternalDependencyError):
            generator.generate_report(
                TEMPLATE, context, OUT / "bad_img.docx",
                image_vars={"chart": str(OUT / "不存在的图.png")},
                template_info=info)

    def test_generate_no_placeholders_fails(self):
        from docx import Document
        no_ph = OUT / "no_placeholder.docx"
        doc = Document()
        doc.add_paragraph("没有任何占位符")
        doc.save(str(no_ph))
        with pytest.raises(DocumentFormatError):
            generator.generate_report(no_ph, {}, OUT / "should_not_exist.docx")


class TestValidator:
    def test_validate_context(self):
        info = template_parser.parse_template(TEMPLATE)
        context = {"实验名称": "x", "姓名": "", "学号": ""}
        issues = validator.validate_context(context, info)
        names = {i.field for i in issues}
        assert "姓名" in names
        # chart 在 image_placeholders 中应被跳过
        issues2 = validator.validate_context(context, info, image_placeholders={"chart"})
        names2 = {i.field for i in issues2}
        assert "chart" not in names2

    def test_validate_mappings(self):
        from formora.core.models import FieldMapping
        mappings = [
            FieldMapping(template_field="姓名", source_type="column", source_column=""),
            FieldMapping(template_field="备注", source_type="constant", constant_value=""),
            FieldMapping(template_field="学号", source_type="constant", constant_value="123"),
        ]
        issues = validator.validate_mappings(mappings)
        by_field = {i.field: i for i in issues}
        assert by_field["姓名"].severity == "warning"
        assert by_field["备注"].severity == "warning"
        assert "学号" not in by_field


class TestChart:
    def test_generate_line_chart(self):
        png = OUT / "line.png"
        chart.generate_line_chart([1, 2, 3], [4, 5, 6], png, title="折线")
        assert png.exists() and png.stat().st_size > 0

    def test_generate_bar_chart(self):
        png = OUT / "bar.png"
        chart.generate_bar_chart(["A", "B"], [1.0, 2.0], png, title="柱状")
        assert png.exists() and png.stat().st_size > 0

    def test_empty_values_raises(self):
        with pytest.raises(DocumentFormatError):
            chart.generate_line_chart([], [], OUT / "empty.png")


class TestErrors:
    def test_unsupported_data_format(self):
        bad = SAMPLES / "x.dat"
        bad.write_bytes(b"data")
        try:
            with pytest.raises(DocumentFormatError):
                data_reader.read_data(bad)
        finally:
            bad.unlink()