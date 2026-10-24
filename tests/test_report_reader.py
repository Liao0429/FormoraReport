"""报告文本读取测试：段落/表格/图片提示抽取。"""

from __future__ import annotations

from pathlib import Path

from formora.services.report_reader import (
    extract_report_text,
    extract_report_text_with_sources,
)


def _make_docx_with_content(path: Path) -> None:
    from docx import Document
    from PIL import Image

    doc = Document()
    doc.add_paragraph("报告正文第一段")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "单元格A"
    table.cell(0, 1).text = "单元格B"
    img_path = path.parent / "mini.png"
    Image.new("RGB", (2, 2), "red").save(str(img_path))
    doc.add_picture(str(img_path))
    doc.save(str(path))


class TestExtractReportText:
    def test_paragraph_and_table(self, tmp_path):
        docx = tmp_path / "r.docx"
        _make_docx_with_content(docx)
        text = extract_report_text(docx)
        assert "报告正文第一段" in text
        assert "[表格 1]" in text
        assert "单元格A | 单元格B" in text

    def test_image_hint(self, tmp_path):
        docx = tmp_path / "r.docx"
        _make_docx_with_content(docx)
        text = extract_report_text(docx)
        assert "图片" in text and "1 张" in text

    def test_with_sources_appended(self, tmp_path):
        docx = tmp_path / "r.docx"
        _make_docx_with_content(docx)
        text = extract_report_text_with_sources(
            docx, ai_sources_text="AI 依据：实验报告模板.docx"
        )
        assert text.endswith("AI 依据：实验报告模板.docx")

    def test_without_sources_no_suffix(self, tmp_path):
        docx = tmp_path / "r.docx"
        _make_docx_with_content(docx)
        text = extract_report_text_with_sources(docx)
        assert "AI 依据" not in text
