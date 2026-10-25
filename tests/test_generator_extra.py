"""报告生成器补充测试：目录创建、解析缓存、依赖缺失、图片包装。"""

from __future__ import annotations

import sys
import threading
from pathlib import Path

import pytest

from formora.core import generator
from formora.core.errors import DocumentFormatError, ExternalDependencyError
from formora.core.models import TemplateInfo


def _make_mini_png(path: Path) -> Path:
    from PIL import Image

    img = Image.new("RGB", (4, 4), "#4a90d9")
    img.save(str(path))
    return path


class TestOutputDirs:
    def test_creates_nested_parent(self, tmp_path):
        """输出目录不存在时应自动创建。"""
        from docx import Document

        tpl = tmp_path / "tpl.docx"
        doc = Document()
        doc.add_paragraph("{{实验名称}}")
        doc.save(str(tpl))
        out = tmp_path / "deep" / "nested" / "out.docx"
        generator.generate_report(tpl, {"实验名称": "x"}, out)
        assert out.exists()

    def test_reuses_template_info(self, tmp_path, monkeypatch):
        """传入 template_info 时不应重复解析模板。"""
        from docx import Document

        tpl = tmp_path / "tpl.docx"
        doc = Document()
        doc.add_paragraph("{{实验名称}}")
        doc.save(str(tpl))
        info = TemplateInfo(path=tpl, placeholders=["实验名称"])

        import formora.core.generator as gen

        monkeypatch.setattr(
            gen, "parse_template", lambda *a, **k: pytest.fail("不应重新解析")
        )
        out = generator.generate_report(
            tpl, {"实验名称": "x"}, tmp_path / "out.docx", template_info=info
        )
        assert out.exists()


class TestDependencyErrors:
    def test_docxtpl_missing_raises(self, tmp_path, monkeypatch):
        from docx import Document

        tpl = tmp_path / "tpl.docx"
        doc = Document()
        doc.add_paragraph("{{实验名称}}")
        doc.save(str(tpl))
        info = TemplateInfo(path=tpl, placeholders=["实验名称"])
        monkeypatch.setitem(sys.modules, "docxtpl", None)
        with pytest.raises(ExternalDependencyError):
            generator.generate_report(
                tpl, {"实验名称": "x"}, tmp_path / "out.docx", template_info=info
            )

    def test_no_placeholders_raises(self, tmp_path):
        from docx import Document

        tpl = tmp_path / "tpl.docx"
        doc = Document()
        doc.add_paragraph("没有任何占位符")
        doc.save(str(tpl))
        info = TemplateInfo(path=tpl, placeholders=[], loop_blocks=[])
        with pytest.raises(DocumentFormatError):
            generator.generate_report(
                tpl, {}, tmp_path / "out.docx", template_info=info
            )


class TestRenderImage:
    def test_missing_image_raises(self, tmp_path):
        from docx import Document

        tpl = tmp_path / "tpl.docx"
        doc = Document()
        doc.add_paragraph("{{chart}}")
        doc.save(str(tpl))
        from docxtpl import DocxTemplate

        t = DocxTemplate(str(tpl))
        with pytest.raises(ExternalDependencyError):
            generator._render_image(t, "chart", str(tmp_path / "不存在.png"))

    def test_existing_image_returns_inline_image(self, tmp_path):
        from docx import Document

        tpl = tmp_path / "tpl.docx"
        doc = Document()
        doc.add_paragraph("{{chart}}")
        doc.save(str(tpl))
        png = _make_mini_png(tmp_path / "img.png")
        from docxtpl import DocxTemplate

        t = DocxTemplate(str(tpl))
        obj = generator._render_image(t, "chart", str(png))
        # InlineImage 对象应有 _width 等属性（docxtpl 内部实现）
        assert obj is not None


class TestTemplateParseCache:
    """模板解析缓存：命中复用、修改后失效、并发安全。"""

    @pytest.fixture(autouse=True)
    def _clear_cache(self):
        generator.clear_template_cache()
        yield
        generator.clear_template_cache()

    @staticmethod
    def _make_template(tmp_path: Path, placeholder: str = "实验名称") -> Path:
        from docx import Document

        tpl = tmp_path / "tpl.docx"
        doc = Document()
        doc.add_paragraph(f"{{{{{placeholder}}}}}")
        doc.save(str(tpl))
        return tpl

    def test_cache_hit_reuses_parsed_skeleton(self, tmp_path, monkeypatch):
        """同一模板连续生成多份时，docx 只解析一次（缓存命中）。"""
        import docx

        tpl = self._make_template(tmp_path)
        info = TemplateInfo(path=tpl, placeholders=["实验名称"])
        real_document = docx.Document
        parses: list[int] = []

        def counting_document(*args, **kwargs):
            parses.append(1)
            return real_document(*args, **kwargs)

        monkeypatch.setattr(docx, "Document", counting_document)

        generator.generate_report(
            tpl, {"实验名称": "A"}, tmp_path / "out1.docx", template_info=info
        )
        generator.generate_report(
            tpl, {"实验名称": "B"}, tmp_path / "out2.docx", template_info=info
        )
        assert len(parses) == 1, "第二次生成应命中缓存，不再解析模板"
        assert (tmp_path / "out1.docx").exists()
        assert (tmp_path / "out2.docx").exists()

    def test_cache_invalidates_after_template_modified(self, tmp_path, monkeypatch):
        """模板文件修改后（mtime/size 变化）应重新解析。"""
        import docx

        tpl = self._make_template(tmp_path)
        info = TemplateInfo(path=tpl, placeholders=["实验名称"])
        real_document = docx.Document
        parses: list[int] = []

        def counting_document(*args, **kwargs):
            parses.append(1)
            return real_document(*args, **kwargs)

        monkeypatch.setattr(docx, "Document", counting_document)

        generator.generate_report(
            tpl, {"实验名称": "A"}, tmp_path / "out1.docx", template_info=info
        )
        assert len(parses) == 1

        # 改写模板内容后再次生成，应触发重新解析
        new_doc = real_document()
        new_doc.add_paragraph("{{实验名称}}（新版本）")
        new_doc.save(str(tpl))

        generator.generate_report(
            tpl, {"实验名称": "B"}, tmp_path / "out2.docx", template_info=info
        )
        assert len(parses) == 2, "模板修改后缓存应失效并重新解析"
        assert (tmp_path / "out2.docx").exists()

    def test_clear_template_cache(self, tmp_path, monkeypatch):
        """clear_template_cache 后首次生成应重新解析。"""
        import docx

        tpl = self._make_template(tmp_path)
        info = TemplateInfo(path=tpl, placeholders=["实验名称"])
        real_document = docx.Document
        parses: list[int] = []

        def counting_document(*args, **kwargs):
            parses.append(1)
            return real_document(*args, **kwargs)

        monkeypatch.setattr(docx, "Document", counting_document)

        generator.generate_report(
            tpl, {"实验名称": "A"}, tmp_path / "out1.docx", template_info=info
        )
        generator.generate_report(
            tpl, {"实验名称": "B"}, tmp_path / "out2.docx", template_info=info
        )
        assert len(parses) == 1

        generator.clear_template_cache()
        generator.generate_report(
            tpl, {"实验名称": "C"}, tmp_path / "out3.docx", template_info=info
        )
        assert len(parses) == 2, "清空缓存后应重新解析"

    def test_concurrent_generation_thread_safe(self, tmp_path):
        """多线程并发生成（模拟多个 QThread 批量任务）应安全且输出正确。"""
        tpl = self._make_template(tmp_path)
        info = TemplateInfo(path=tpl, placeholders=["实验名称"])
        out_dir = tmp_path / "out"
        out_dir.mkdir()
        errors: list[Exception] = []

        def worker(n: int) -> None:
            try:
                for i in range(5):
                    generator.generate_report(
                        tpl,
                        {"实验名称": f"任务{n}-{i}"},
                        out_dir / f"{n}_{i}.docx",
                        template_info=info,
                    )
            except Exception as e:  # pragma: no cover - 失败时记录断言
                errors.append(e)

        threads = [threading.Thread(target=worker, args=(n,)) for n in range(4)]
        for t in threads:
            t.start()
        for t in threads:
            t.join()

        assert not errors
        from docx import Document

        files = sorted(out_dir.glob("*.docx"))
        assert len(files) == 20
        texts = " ".join(p.text for p in Document(str(files[0])).paragraphs)
        assert "任务0-0" in texts


class TestExceptionPaths:
    """异常路径：图片包装/模板加载/渲染/保存失败统一转为 DocumentFormatError。"""

    @pytest.fixture(autouse=True)
    def _clear_cache(self):
        generator.clear_template_cache()
        yield
        generator.clear_template_cache()

    @staticmethod
    def _make_template(tmp_path: Path, placeholder: str = "实验名称") -> Path:
        from docx import Document

        tpl = tmp_path / "tpl.docx"
        doc = Document()
        doc.add_paragraph(f"{{{{{placeholder}}}}}")
        doc.save(str(tpl))
        return tpl

    def test_render_image_failure_raises(self, tmp_path, monkeypatch):
        """InlineImage 构造抛异常时，_render_image 应包装为 DocumentFormatError。"""
        import docxtpl
        from docxtpl import DocxTemplate

        tpl = self._make_template(tmp_path, placeholder="chart")
        png = _make_mini_png(tmp_path / "img.png")
        t = DocxTemplate(str(tpl))

        def _boom(*args, **kwargs):
            raise RuntimeError("image boom")

        monkeypatch.setattr(docxtpl, "InlineImage", _boom)
        with pytest.raises(DocumentFormatError, match="插入图片"):
            generator._render_image(t, "chart", str(png))

    def test_load_template_failure_raises(self, tmp_path, monkeypatch):
        """DocxTemplate 构造失败时，generate_report 应包装为 DocumentFormatError。"""
        import docxtpl

        tpl = self._make_template(tmp_path)
        info = TemplateInfo(path=tpl, placeholders=["实验名称"])

        def _boom(*args, **kwargs):
            raise RuntimeError("load boom")

        monkeypatch.setattr(docxtpl, "DocxTemplate", _boom)
        with pytest.raises(DocumentFormatError, match="无法加载模板"):
            generator.generate_report(
                tpl, {"实验名称": "x"}, tmp_path / "out.docx", template_info=info
            )

    def test_render_failure_raises(self, tmp_path, monkeypatch):
        """tpl.render 抛异常时，generate_report 应包装为 DocumentFormatError。"""
        tpl = self._make_template(tmp_path)
        info = TemplateInfo(path=tpl, placeholders=["实验名称"])

        class _BrokenRenderTpl:
            def render(self, *args, **kwargs):
                raise RuntimeError("render boom")

            def save(self, path):
                pass

        monkeypatch.setattr(
            generator, "_get_cached_template", lambda *a, **k: _BrokenRenderTpl()
        )
        with pytest.raises(DocumentFormatError, match="模板渲染失败"):
            generator.generate_report(
                tpl, {"实验名称": "x"}, tmp_path / "out.docx", template_info=info
            )

    def test_save_failure_raises(self, tmp_path, monkeypatch):
        """tpl.save 抛异常时，generate_report 应包装为 DocumentFormatError。"""
        tpl = self._make_template(tmp_path)
        info = TemplateInfo(path=tpl, placeholders=["实验名称"])

        class _BrokenSaveTpl:
            def render(self, *args, **kwargs):
                pass

            def save(self, path):
                raise RuntimeError("save boom")

        monkeypatch.setattr(
            generator, "_get_cached_template", lambda *a, **k: _BrokenSaveTpl()
        )
        with pytest.raises(DocumentFormatError, match="保存输出文档失败"):
            generator.generate_report(
                tpl, {"实验名称": "x"}, tmp_path / "out.docx", template_info=info
            )
